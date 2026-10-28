import io
import os
from datetime import datetime
import streamlit as st
from docx import Document
import openpyxl

# Configuración visual de la aplicación web
st.set_page_config(
    page_title="Generador de Contratos - Carga Directa",
    page_icon="📄",
    layout="centered"
)

def pct(valor):
    """Convierte un número flotante (ej: 0.05) a formato porcentaje (5.00%)"""
    try:
        return f"{(float(valor) * 100):.2f}%"
    except:
        return "0.00%"

def procesar_contrato(id_contrato, archivo_excel_subido, archivo_plantilla_subida):
    """Lee el Excel cargado por el usuario, busca el ID y genera el archivo Word en memoria"""
    try:
        # 1. Cargar el archivo Excel en memoria utilizando openpyxl
        wb = openpyxl.load_workbook(archivo_excel_subido, data_only=True)
        
        # Validación de seguridad para asegurar que exista la pestaña requerida
        if "Actual" not in wb.sheetnames:
            st.error("❌ El archivo de Excel no contiene una pestaña llamada 'Actual'.")
            return None, None
            
        ws_actual = wb["Actual"]
        
        # 2. Buscar la fila mediante coincidencia exacta con el ID (Columna A / Columna 1)
        fila_encontrada = None
        for row in range(1, ws_actual.max_row + 1):
            if str(ws_actual.cell(row=row, column=1).value).strip() == str(id_contrato).strip():
                fila_encontrada = row
                break
                
        if not fila_encontrada:
            st.error(f"❌ No se encontró el ID '{id_contrato}' en la columna A de la hoja 'Actual'.")
            return None, None

        # 3. Extraer los datos de la fila (Mapeo idéntico al código original)
        account_name = str(ws_actual.cell(row=fila_encontrada, column=3).value or '').strip()
        dot          = str(ws_actual.cell(row=fila_encontrada, column=6).value or '').strip()
        address      = str(ws_actual.cell(row=fila_encontrada, column=12).value or '').strip()
        legal_rep    = str(ws_actual.cell(row=fila_encontrada, column=13).value or '').strip()
        client_email = str(ws_actual.cell(row=fila_encontrada, column=14).value or '').strip()
        
        # Tratamiento especial de la celda de tipo Fecha
        raw_date = ws_actual.cell(row=fila_encontrada, column=15).value
        if isinstance(raw_date, datetime):
            fecha_str = raw_date.strftime("%m/%d/%Y")
        else:
            fecha_str = str(raw_date or '')

        mc           = str(ws_actual.cell(row=fila_encontrada, column=16).value or '').strip()
        
        # Extracción de valores numéricos para porcentajes
        fee_val      = ws_actual.cell(row=fila_encontrada, column=17).value or 0
        reserve_val  = ws_actual.cell(row=fila_encontrada, column=18).value or 0
        
        # Cálculo de la variable idéntico a la lógica matemática previa
        try:
            variable_val = 1 - float(fee_val) - float(reserve_val)
        except:
            variable_val = 0
            
        fee      = pct(fee_val)
        reserve  = pct(reserve_val)
        variable = pct(variable_val)
        
        bank_name   = str(ws_actual.cell(row=fila_encontrada, column=19).value or '').strip()
        wire_number = str(ws_actual.cell(row=fila_encontrada, column=20).value or '').strip()
        acc_number  = str(ws_actual.cell(row=fila_encontrada, column=21).value or '').strip()

        # Diccionario estructurado de mapeo para las etiquetas de reemplazo en Word
        reemplazos = {
            "{{Account Name}}": account_name,
            "{{Legal Representative}}": legal_rep,
            "{{Fee}}": fee,
            "{{Reserve}}": reserve,
            "{{Variable}}": variable,
            "{{Date}}": fecha_str,
            "{{DOT}}": dot,
            "{{MC}}": mc,
            "{{Address}}": address,
            "{{Client Email Address}}": client_email,
            "{{Bank Name}}": bank_name,
            "{{Account Number}}": acc_number,
            "{{Wire Number}}": wire_number
        }

        # 4. Abrir la plantilla de Word y ejecutar los reemplazos de texto
        doc = Document(archivo_plantilla_subida)
        
        # Reemplazar texto en los párrafos del cuerpo del documento
        for p in doc.paragraphs:
            for placeholder, valor in reemplazos.items():
                if placeholder in p.text:
                    p.text = p.text.replace(placeholder, valor)
                    
        # Reemplazar texto dentro de las celdas de las tablas del documento
        for tabla in doc.tables:
            for fila_tabla in tabla.rows:
                for celda in fila_tabla.cells:
                    for p in celda.paragraphs:
                        for placeholder, valor in reemplazos.items():
                            if placeholder in p.text:
                                p.text = p.text.replace(placeholder, valor)

        # Configuración limpia del nombre del archivo final de salida
        fecha_limpia = fecha_str.replace("/", "_")
        nombre_descarga = f"{fecha_limpia} - FACTORING CONTRACT - {account_name}.docx"
        
        return doc, nombre_descarga

    except Exception as e:
        st.error(f"⚠️ Ocurrió un error inesperado al procesar los archivos: {str(e)}")
        return None, None


# --- INTERFAZ DEL USUARIO (STREAMLIT) ---
st.title("📄 Generador Inmediato de Contratos")
st.write("Sube los archivos fuentes requeridos, digita el identificador y descarga el resultado de forma instantánea.")

# Componentes de arrastre de archivos en la barra principal
excel_file = st.file_uploader("1. Sube tu base de datos de Excel (.xlsx)", type=["xlsx"])
word_file = st.file_uploader("2. Sube la plantilla base de Word (.docx)", type=["docx"])

st.markdown("---")

# Control de captura de texto para el identificador único del contrato
id_input = st.text_input("3. Ingresa el ID del Contrato (Columna A):")

if excel_file and word_file and id_input:
    # Procesar de forma automática y transparente al completarse las entradas
    with st.spinner("Buscando ID y estructurando el contrato..."):
        documento_final, nombre_salida = procesar_contrato(id_input, excel_file, word_file)
    
    if documento_final and nombre_salida:
        st.success(f"✨ ¡Contrato estructurado con éxito para el cliente!")
        
        # Canalizar el archivo a un stream binario en memoria para evitar almacenamiento en disco local
        buffer_binario = io.BytesIO()
        documento_final.save(buffer_binario)
        buffer_binario.seek(0)
        
        # Disparador nativo del navegador del usuario para iniciar la descarga del documento
        st.download_button(
            label="⬇️ Descargar Contrato (.docx)",
            data=buffer_binario.getvalue(),
            file_name=nombre_salida,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
else:
    st.info("💡 Para comenzar, carga el archivo Excel, la plantilla de Word e ingresa el ID del contrato.")
